"""Read-only corpus loading and lightweight, auditable cleaning."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable
import csv
import json
import re


@dataclass
class Document:
    doc_id: str
    text: str
    source: str
    source_row: int | None = None


@dataclass
class Corpus:
    documents: list[Document]
    cleaning_log: list[dict[str, Any]] = field(default_factory=list)
    stats: dict[str, int] = field(default_factory=dict)

    @property
    def ids(self) -> list[str]:
        return [d.doc_id for d in self.documents]

    @property
    def texts(self) -> list[str]:
        return [d.text for d in self.documents]


ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "big5", "utf-16")


def _decode_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    raise UnicodeError(f"无法识别文本编码：{path.name}。支持 UTF-8、GB18030、Big5、UTF-16。")


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        # pandas/numpy NA without importing either here
        if value != value:
            return ""
    except Exception:
        pass
    return str(value)


def _unique_id(candidate: str, seen: set[str]) -> str:
    base = candidate.strip() or "document"
    if base not in seen:
        seen.add(base)
        return base
    suffix = 2
    while f"{base}#{suffix}" in seen:
        suffix += 1
    result = f"{base}#{suffix}"
    seen.add(result)
    return result


def split_text(text: str, mode: str) -> list[str]:
    """Split pasted or plain-text data as document/file, line, or paragraph."""
    if mode == "整篇/每文件":
        return [text]
    if mode == "每行":
        return text.splitlines()
    if mode == "每段":
        return re.split(r"\n\s*\n+", text)
    raise ValueError(f"未知文本切分方式：{mode}")


def _table_records(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        last_error: Exception | None = None
        for enc in ENCODINGS:
            try:
                with path.open("r", encoding=enc, newline="") as handle:
                    sample = handle.read(8192)
                    handle.seek(0)
                    delimiter = "\t" if suffix == ".tsv" else None
                    if delimiter is None:
                        try:
                            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
                        except csv.Error:
                            delimiter = ","
                    reader = csv.DictReader(handle, delimiter=delimiter)
                    if not reader.fieldnames:
                        return [], []
                    return [dict(row) for row in reader], [str(c) for c in reader.fieldnames]
            except UnicodeDecodeError as exc:
                last_error = exc
        raise UnicodeError(f"无法识别表格编码：{path.name}") from last_error
    if suffix == ".xlsx":
        import pandas as pd
        frame = pd.read_excel(path, dtype=object)
        cols = [str(c) for c in frame.columns]
        frame.columns = cols
        return frame.to_dict(orient="records"), cols
    if suffix == ".jsonl":
        records: list[dict[str, Any]] = []
        with path.open("r", encoding="utf-8-sig") as handle:
            for line_no, line in enumerate(handle, 1):
                if not line.strip():
                    continue
                try:
                    value = json.loads(line)
                except json.JSONDecodeError as exc:
                    raise ValueError(f"JSONL 第 {line_no} 行格式错误：{exc.msg}") from exc
                records.append(value if isinstance(value, dict) else {"text": value})
        cols = list(dict.fromkeys(str(k) for row in records for k in row))
        return records, cols
    if suffix == ".json":
        value = json.loads(_decode_text(path))
        if isinstance(value, list):
            records = [row if isinstance(row, dict) else {"text": row} for row in value]
        elif isinstance(value, dict):
            records = value.get("data", value.get("records", [value]))
            if not isinstance(records, list):
                records = [records]
            records = [row if isinstance(row, dict) else {"text": row} for row in records]
        else:
            records = [{"text": value}]
        cols = list(dict.fromkeys(str(k) for row in records for k in row))
        return records, cols
    raise ValueError(f"不支持的文件类型：{path.suffix}。可读 TXT、CSV、TSV、XLSX、JSONL、JSON。")


def table_columns(path: str | Path) -> list[str]:
    """Return column names for UI column selectors."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        last_error: Exception | None = None
        for enc in ENCODINGS:
            try:
                with p.open("r", encoding=enc, newline="") as handle:
                    sample = handle.read(8192)
                    handle.seek(0)
                    delimiter = "\t" if suffix == ".tsv" else None
                    if delimiter is None:
                        try:
                            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
                        except csv.Error:
                            delimiter = ","
                    return [str(x) for x in (next(csv.reader(handle, delimiter=delimiter), []))]
            except UnicodeDecodeError as exc:
                last_error = exc
        raise UnicodeError(f"无法识别表格编码：{p.name}") from last_error
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(p, read_only=True, data_only=True)
        try:
            row = next(book.active.iter_rows(min_row=1, max_row=1, values_only=True), ())
            return [str(x) for x in row if x is not None]
        finally:
            book.close()
    if suffix == ".jsonl":
        with p.open("r", encoding="utf-8-sig") as handle:
            for line in handle:
                if line.strip():
                    value = json.loads(line)
                    return list(map(str, value.keys())) if isinstance(value, dict) else ["text"]
        return []
    if suffix == ".json":
        value = json.loads(_decode_text(p))
        if isinstance(value, dict):
            value = value.get("data", value.get("records", [value]))
        if isinstance(value, list) and value and isinstance(value[0], dict):
            return [str(k) for k in value[0].keys()]
        return ["text"] if value is not None else []
    raise ValueError(f"不支持的文件类型：{p.suffix}")


def load_corpus(
    paths: Iterable[str | Path] = (),
    *,
    pasted_text: str = "",
    text_column: str = "",
    id_column: str = "",
    split_mode: str = "整篇/每文件",
    deduplicate: bool = False,
    cancel_event: Any = None,
    progress: Any = None,
) -> Corpus:
    """Load supported files without modifying them; empty rows are excluded."""
    candidates: list[tuple[str, str, int | None, str]] = []
    log: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    input_paths = [Path(p) for p in paths]
    expanded: list[Path] = []
    directory_identity: dict[Path, str] = {}
    for path in input_paths:
        if path.is_dir():
            members = sorted((p for p in path.rglob("*.txt") if p.is_file()), key=lambda p: p.relative_to(path).as_posix().casefold())
            expanded.extend(members)
            directory_identity.update({p: p.relative_to(path).as_posix() for p in members})
        else:
            expanded.append(path)
    if pasted_text:
        parts = split_text(pasted_text, split_mode)
        for i, text in enumerate(parts, 1):
            candidates.append((f"粘贴文本:{i:06d}" if split_mode != "整篇/每文件" else "粘贴文本", text, i, "粘贴文本"))
        log.append({"来源": "粘贴文本", "读入行/段": len(parts)})

    for file_no, path in enumerate(expanded, 1):
        if cancel_event is not None and cancel_event.is_set():
            raise InterruptedError("已取消导入")
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(f"找不到文件：{path}")
        source = path.name
        before = len(candidates)
        if path.suffix.lower() == ".txt":
            text = _decode_text(path)
            parts = split_text(text, split_mode)
            identity_source = directory_identity.get(path, source)
            for i, part in enumerate(parts, 1):
                identity = identity_source if split_mode == "整篇/每文件" else f"{identity_source}:{i:06d}"
                candidates.append((identity, part, i, str(path)))
            rows_read = len(parts)
        else:
            records, columns = _table_records(path)
            if text_column and text_column not in columns:
                raise ValueError(f"{source} 不包含指定的文本列“{text_column}”。可用列：{', '.join(columns)}")
            if id_column and id_column not in columns:
                raise ValueError(f"{source} 不包含指定的 ID 列“{id_column}”。可用列：{', '.join(columns)}")
            chosen_text = text_column if text_column in columns else ""
            if not chosen_text:
                # An explicit field is preferred; otherwise choose a reasonable first text-like column.
                chosen_text = next((c for c in columns if c.casefold() in {"text", "content", "文本", "正文", "内容"}), "")
                if not chosen_text:
                    chosen_text = next((c for c in columns if c != id_column), "")
            if not chosen_text:
                raise ValueError(f"{source} 没有可用文本列，请指定文本列。列：{', '.join(columns)}")
            for i, row in enumerate(records, 1):
                text = _as_text(row.get(chosen_text, ""))
                raw_id = _as_text(row.get(id_column, "")).strip() if id_column and id_column in row else ""
                identity = raw_id or f"{source}:{i:06d}"
                candidates.append((identity, text, i, str(path)))
            rows_read = len(records)
        added = candidates[before:]
        source_log = {"来源": str(path), "读入行/段": rows_read}
        if path.suffix.lower() != ".txt":
            source_log["文本列"] = chosen_text
        log.append(source_log)
        if progress:
            progress(file_no, len(expanded), f"已读取 {source}")

    docs: list[Document] = []
    total_empty = 0
    total_dupes = 0
    dedupe_seen: set[str] = set()
    per_source = {row["来源"]: row for row in log}
    for candidate_id, text, row_no, source in candidates:
        text = _as_text(text)
        if not text.strip():
            total_empty += 1
            if source in per_source:
                per_source[source]["空值排除"] = per_source[source].get("空值排除", 0) + 1
            continue
        # Keep original content; normalize whitespace only for optional exact-text duplicate detection.
        key = re.sub(r"\s+", " ", text).strip()
        if deduplicate and key in dedupe_seen:
            total_dupes += 1
            if source in per_source:
                per_source[source]["去重排除"] = per_source[source].get("去重排除", 0) + 1
            continue
        dedupe_seen.add(key)
        docs.append(Document(_unique_id(candidate_id, seen_ids), text, source, row_no))
    stats = {"读入记录": len(candidates), "空值排除": total_empty, "重复排除": total_dupes, "保留文档": len(docs)}
    if not docs:
        raise ValueError("没有可分析的非空文本。请检查输入文件、文本列或切分方式。")
    return Corpus(documents=docs, cleaning_log=log, stats=stats)

